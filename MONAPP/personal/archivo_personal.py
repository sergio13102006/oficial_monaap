from io import BytesIO
from datetime import datetime

from django.contrib.staticfiles import finders
from django.http import HttpResponse
from django.utils.text import slugify

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image as RLImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


PAGE_WIDTH, PAGE_HEIGHT = A4


def _nombre_archivo_personal(extension):
    fecha = datetime.now().strftime("%Y%m%d")
    return f"archivo_personal_monakeratina_{slugify(fecha)}.{extension}"


def _find_first_static(paths):
    for path in paths:
        found = finders.find(path)
        if found:
            return found
    return None


def _get_logo_path():
    return _find_first_static([
        "compras/img/logo_monakeratina.png",
        "core/img/logo_monakeratina.png",
        "core/img/logo_monakeratina.jpg",
    ])


def _get_watermark_path():
    return _find_first_static([
        "compras/img/logo_monakeratina_watermark.png",
        "core/img/logo_monakeratina_watermark.png",
        "core/img/logo_monakeratina.png",
    ])


def _draw_watermark(canvas, doc):
    watermark_path = _get_watermark_path()
    if not watermark_path:
        return

    canvas.saveState()

    try:
        if hasattr(canvas, "setFillAlpha"):
            canvas.setFillAlpha(0.06)
    except Exception:
        pass

    image_width = 18 * mm
    image_height = 18 * mm
    gap_x = 16 * mm
    gap_y = 18 * mm

    x = 10 * mm
    while x < PAGE_WIDTH:
        y = 12 * mm
        while y < PAGE_HEIGHT:
            canvas.drawImage(
                watermark_path,
                x,
                y,
                width=image_width,
                height=image_height,
                preserveAspectRatio=True,
                mask="auto",
            )
            y += image_height + gap_y
        x += image_width + gap_x

    canvas.restoreState()


def build_personal_pdf_response(personal_list):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#000000"),
        alignment=1,
        spaceAfter=2,
    )
    small_label_style = ParagraphStyle(
        "SmallLabelStyle",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=9,
        textColor=colors.HexColor("#333333"),
    )
    small_value_style = ParagraphStyle(
        "SmallValueStyle",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=13,
        textColor=colors.HexColor("#111111"),
    )
    footer_style = ParagraphStyle(
        "FooterStyle",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#666666"),
        alignment=1,
    )

    elements = []

    logo_path = _get_logo_path()
    logo_flowable = ""
    if logo_path:
        try:
            logo_flowable = RLImage(logo_path, width=34 * mm, height=16 * mm)
        except Exception:
            logo_flowable = ""

    header_table = Table(
        [[logo_flowable, Paragraph("Archivo de personal", title_style), ""]],
        colWidths=[42 * mm, 100 * mm, 26 * mm],
    )
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, 0), "LEFT"),
        ("ALIGN", (1, 0), (1, 0), "CENTER"),
        ("LINEBELOW", (0, 0), (-1, 0), 1.2, colors.black),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 8))

    top_info = Table(
        [[
            Table([[Paragraph("NUMERO", small_label_style)], [Paragraph(f"{len(personal_list):04d}", small_value_style)]], colWidths=[28 * mm]),
            Table([
                [Paragraph("DE", small_label_style)],
                [Paragraph("Monakeratina", small_value_style)],
                [Paragraph("Sistema de personal", styles["BodyText"])],
                [Paragraph("Archivo interno", styles["BodyText"])],
            ], colWidths=[64 * mm]),
            Table([
                [Paragraph("PARA", small_label_style)],
                [Paragraph("Personal registrado", small_value_style)],
                [Paragraph("Base de personal de MONAAPP", styles["BodyText"])],
            ], colWidths=[74 * mm]),
        ]],
        colWidths=[30 * mm, 66 * mm, 76 * mm],
    )
    top_info.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEABOVE", (0, 0), (-1, -1), 1.0, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(top_info)
    elements.append(Spacer(1, 10))

    meta_table = Table(
        [[
            Table([[Paragraph("FECHA", small_label_style)], [Paragraph(datetime.now().strftime("%d/%m/%Y"), small_value_style)]], colWidths=[28 * mm]),
            Table([[Paragraph("MODULO", small_label_style)], [Paragraph("Personal", small_value_style)]], colWidths=[28 * mm]),
            Table([[Paragraph("ESTADO", small_label_style)], [Paragraph("Generado", small_value_style)]], colWidths=[28 * mm]),
        ]],
        colWidths=[30 * mm, 30 * mm, 30 * mm],
    )
    meta_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEABOVE", (0, 0), (-1, -1), 1.0, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 10))

    table_data = [["Documento", "Nombre", "Telefono", "Correo", "Rol", "Estado"]]
    if personal_list:
        for persona in personal_list:
            table_data.append([
                f"{persona.tipo_documento} {persona.numero_documento}",
                f"{persona.nombres} {persona.apellidos}",
                str(persona.telefono or "-"),
                str(persona.correo or "-"),
                str(persona.rol or "-"),
                "Activo" if persona.activo else "Inactivo",
            ])
    else:
        table_data.append(["No hay personal registrado.", "", "", "", "", ""])

    detail_table = Table(
        table_data,
        colWidths=[32 * mm, 42 * mm, 24 * mm, 40 * mm, 20 * mm, 16 * mm],
        repeatRows=1,
    )
    detail_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 1.0, colors.black),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#777777")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 1), (-1, -1), 0.35, colors.HexColor("#CFCFCF")),
    ]))
    elements.append(detail_table)
    elements.append(Spacer(1, 6))

    total_table = Table(
        [["Total registros", str(len(personal_list))]],
        colWidths=[34 * mm, 34 * mm],
        hAlign="RIGHT",
    )
    total_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 1.0, colors.black),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (0, 0), (0, 0), "LEFT"),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(total_table)
    elements.append(Spacer(1, 10))
    elements.append(Paragraph("Generado por MonaApp / Monakeratina", footer_style))

    doc.build(elements, onFirstPage=_draw_watermark, onLaterPages=_draw_watermark)

    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{_nombre_archivo_personal("pdf")}"'
    return response


def build_personal_excel_response(personal_list):
    wb = Workbook()
    ws = wb.active
    ws.title = "Personal"

    def apply_style(cell, font=None, fill=None, border=None, alignment=None):
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if border:
            cell.border = border
        if alignment:
            cell.alignment = alignment

    def style_range(start_row, end_row, start_col, end_col, **styles):
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                apply_style(ws.cell(row=row, column=col), **styles)

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B12"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 16

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[11].height = 22
    ws.row_dimensions[12].height = 22

    color_dark = "3B261A"
    color_brown = "A67C52"
    color_soft = "F7F0E8"
    color_soft_2 = "FBF7F2"
    color_line = "D8C3B5"
    color_text = "2F241D"
    color_white = "FFFFFF"

    title_font = Font(bold=True, size=16, color=color_text)
    section_font = Font(bold=True, size=11, color=color_white)
    label_font = Font(bold=True, size=10, color=color_dark)
    value_font = Font(size=10, color=color_text)
    table_header_font = Font(bold=True, size=10, color=color_white)
    footer_font = Font(italic=True, size=9, color=color_brown)

    dark_fill = PatternFill("solid", fgColor=color_dark)
    brown_fill = PatternFill("solid", fgColor=color_brown)
    soft_fill = PatternFill("solid", fgColor=color_soft)
    soft_fill_2 = PatternFill("solid", fgColor=color_soft_2)

    thin_side = Side(style="thin", color=color_line)
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    logo_path = _get_logo_path()
    if logo_path:
        try:
            img = XLImage(logo_path)
            img.width = 145
            img.height = 58
            ws.add_image(img, "B1")
        except Exception:
            pass

    style_range(1, 2, 2, 6, border=thin_border, alignment=center)
    ws.merge_cells("B1:F2")
    ws["B1"] = "Archivo de Personal"
    apply_style(ws["B1"], font=title_font, alignment=center)

    style_range(4, 4, 2, 6, fill=dark_fill, border=thin_border, alignment=left)
    ws.merge_cells("B4:F4")
    ws["B4"] = "DATOS GENERALES"
    apply_style(ws["B4"], font=section_font, alignment=left)

    info_rows = [
        (5, "Fecha", datetime.now().strftime("%d/%m/%Y")),
        (6, "Modulo", "Personal"),
        (7, "Estado", "Generado"),
        (8, "Total de registros", len(personal_list)),
    ]
    for row_num, label, value in info_rows:
        ws[f"B{row_num}"] = label
        apply_style(ws[f"B{row_num}"], font=label_font, fill=soft_fill, border=thin_border, alignment=left)
        style_range(row_num, row_num, 3, 6, fill=soft_fill_2, border=thin_border, alignment=left)
        ws.merge_cells(start_row=row_num, start_column=3, end_row=row_num, end_column=6)
        ws.cell(row=row_num, column=3, value=value)
        apply_style(ws.cell(row=row_num, column=3), font=value_font, fill=soft_fill_2, border=thin_border, alignment=left)

    style_range(11, 11, 2, 7, fill=dark_fill, border=thin_border, alignment=left)
    ws.merge_cells("B11:G11")
    ws["B11"] = "DETALLE DE PERSONAL"
    apply_style(ws["B11"], font=section_font, alignment=left)

    header_row = 12
    headers = ["Documento", "Nombre", "Telefono", "Correo", "Rol", "Estado"]
    for col_num, header in enumerate(headers, start=2):
        cell = ws.cell(row=header_row, column=col_num, value=header)
        apply_style(cell, font=table_header_font, fill=brown_fill, border=thin_border, alignment=center)

    current_row = 13
    if personal_list:
        for persona in personal_list:
            values = [
                f"{persona.tipo_documento} {persona.numero_documento}",
                f"{persona.nombres} {persona.apellidos}",
                persona.telefono or "-",
                persona.correo or "-",
                persona.rol or "-",
                "Activo" if persona.activo else "Inactivo",
            ]
            row_fill = soft_fill if current_row % 2 == 0 else soft_fill_2
            for offset, value in enumerate(values, start=2):
                cell = ws.cell(row=current_row, column=offset, value=value)
                apply_style(cell, font=value_font, fill=row_fill, border=thin_border, alignment=left)
            current_row += 1
    else:
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=7)
        ws.cell(row=current_row, column=2, value="No hay personal registrado.")
        style_range(current_row, current_row, 2, 7, fill=soft_fill_2, border=thin_border, alignment=center)

    total_row = current_row + 1
    apply_style(
        ws.cell(row=total_row, column=6, value="TOTAL"),
        font=label_font,
        fill=soft_fill,
        border=thin_border,
        alignment=right,
    )
    apply_style(
        ws.cell(row=total_row, column=7, value=len(personal_list)),
        font=label_font,
        fill=soft_fill,
        border=thin_border,
        alignment=right,
    )

    footer_row = total_row + 2
    style_range(footer_row, footer_row, 2, 7, alignment=center)
    ws.merge_cells(start_row=footer_row, start_column=2, end_row=footer_row, end_column=7)
    ws.cell(row=footer_row, column=2, value="Generado por MonaApp / Monakeratina")
    apply_style(ws.cell(row=footer_row, column=2), font=footer_font, alignment=center)

    ws.print_area = f"B1:G{footer_row}"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{_nombre_archivo_personal("xlsx")}"'
    wb.save(response)
    return response
