from django.shortcuts import render, redirect, get_object_or_404
from django.db.models import Q, Sum, Value, DecimalField
from django.db.models.functions import Coalesce
from django.urls import reverse
from .models import Venta, DetalleVenta, DevolucionVenta, DetalleDevolucion
from inventario.models import Stock
from compras.models import DetalleCompra
from django.contrib import messages
from .forms import VentaForm
from Productos.models import Producto
from django.core.exceptions import ValidationError
import json
from django.db import transaction
from servicios.models import Servicio
from django.http import JsonResponse
from decimal import Decimal
from django.views.decorators.http import require_POST
from django.template.loader import render_to_string
from django.apps import apps
from django.db.models.functions import TruncDate
from datetime import datetime
from io import BytesIO
from django.http import HttpResponse
from django.utils.dateparse import parse_date
from django.contrib.staticfiles import finders
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Spacer, Paragraph, Image as RLImage
import plotly.graph_objects as go

Personal = apps.get_model("personal", "Personal")


def _build_plotly_chart_image(labels, values, tipo="bar", title="", color="#A67C52"):
    """
    Genera una imagen PNG de Plotly para incrustarla en el PDF.

    Si Kaleido/Chrome no está disponible, devuelve None para no romper el reporte.
    """
    if not labels:
        return None

    try:
        fig = go.Figure()

        if tipo == "line":
            fig.add_trace(
                go.Scatter(
                    x=labels,
                    y=values,
                    mode="lines+markers",
                    line=dict(color=color, width=2),
                    marker=dict(color=color, size=8),
                    hovertemplate="%{x}<br>Total: %{y}<extra></extra>",
                )
            )
        else:
            fig.add_trace(
                go.Bar(
                    x=labels,
                    y=values,
                    marker_color=color,
                    hovertemplate="%{x}<br>Total: %{y}<extra></extra>",
                )
            )

        fig.update_layout(
            title=title,
            xaxis_title="Fecha",
            yaxis_title="Total ($)",
            template="plotly_white",
            margin=dict(l=30, r=20, t=50, b=90),
            height=350,
            width=900,
            showlegend=False,
            font=dict(size=12),
        )
        fig.update_xaxes(tickangle=-45)

        image_bytes = fig.to_image(format="png", width=900, height=350, scale=2)
        image_buffer = BytesIO(image_bytes)
        image_buffer.seek(0)
        return image_buffer
    except Exception:
        return None


def _texto_seguro(valor):
    valor = (valor or "").strip()
    return bool(valor) and all(ch.isalnum() or ch.isspace() for ch in valor)


def _find_first_static(paths):
    for path in paths:
        found = finders.find(path)
        if found:
            return found
    return None


def _get_brand_logo_path():
    return _find_first_static(
        [
            "compras/img/logo_monakeratina.png",
            "compras/img/logo_monakeratina.webp",
            "compras/img/logo_monakeratina.jpg",
            "compras/img/logo_monakefratina.png",
            "compras/img/logo_monakefratina.webp",
            "compras/img/logo_monakefratina.jpg",
            "core/img/logo_monakeratina.png",
            "core/img/logo_monakefratina.png",
        ]
    )


def _get_brand_watermark_path():
    return _find_first_static(
        [
            "compras/img/logo_monakeratina_watermark.png",
            "compras/img/logo_monakefratina_watermark.png",
            "compras/img/logo_monakeratina.png",
            "compras/img/logo_monakefratina.png",
            "core/img/logo_monakeratina_watermark.png",
            "core/img/logo_monakefratina_watermark.png",
            "core/img/logo_monakeratina.png",
            "core/img/logo_monakefratina.png",
        ]
    )


def _fmt_money_excel(value):
    try:
        return float(value or 0)
    except Exception:
        return 0.0


def _fmt_money_text(value):
    try:
        num = float(value or 0)
        txt = f"{num:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"$ {txt}"
    except Exception:
        return "$ 0,00"


def _ajustar_stock_producto(producto, delta):
    if delta == 0:
        return Stock.objects.filter(producto=producto).first()

    stock_obj, _ = Stock.objects.get_or_create(
        producto=producto,
        defaults={"cantidad_actual": 0},
    )
    stock_obj = Stock.objects.select_for_update().get(pk=stock_obj.pk)

    stock_anterior = stock_obj.cantidad_actual or 0
    stock_posterior = stock_anterior + delta

    if stock_posterior < 0:
        raise ValidationError(
            f"El stock no puede quedar negativo para {producto.nombre}. "
            f"Actual: {stock_anterior}, movimiento: {delta}."
        )

    stock_obj.cantidad_actual = stock_posterior
    stock_obj.save(update_fields=["cantidad_actual"])
    return stock_obj


def _draw_reporte_watermark(canvas, doc):
    watermark_path = _get_brand_watermark_path()
    if not watermark_path:
        return

    canvas.saveState()

    try:
        if hasattr(canvas, "setFillAlpha"):
            canvas.setFillAlpha(0.06)
    except Exception:
        pass

    page_w, page_h = doc.pagesize

    image_width = 18 * mm
    image_height = 18 * mm
    gap_x = 16 * mm
    gap_y = 18 * mm

    x = 10 * mm
    while x < page_w:
        y = 12 * mm
        while y < page_h:
            canvas.drawImage(
                watermark_path,
                x,
                y,
                width=image_width,
                height=image_height,
                preserveAspectRatio=True,
                mask="auto",
            )
            y += image_height + gap_y
        x += image_width + gap_x

    canvas.restoreState()


def lista_ventas(request):
    q = request.GET.get("q", "").strip()
    estado = request.GET.get("estado", "activa").strip()
    sort = request.GET.get("sort", "fecha").strip()
    direction = request.GET.get("dir", "desc").strip()

    ventas = (
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__producto", "detalles__servicio")
        .annotate(
            total_orden=Coalesce(
                Sum("detalles__subtotal"),
                Value(0),
                output_field=DecimalField(max_digits=12, decimal_places=2),
            )
        )
        .all()
    )

    if q:
        ventas = ventas.filter(
            Q(codigo_venta__icontains=q)
            | Q(cliente__nombre__icontains=q)
            | Q(cliente__apellido__icontains=q)
        )

    if estado:
        ventas = ventas.filter(estado=estado)

    # ==================== ORDENAMIENTO ====================
    if direction not in ["asc", "desc"]:
        direction = "desc"

    if sort == "codigo_venta":
        ventas = ventas.order_by(
            "codigo_venta" if direction == "asc" else "-codigo_venta"
        )

    elif sort == "cliente":
        ventas = ventas.order_by(
            "cliente__nombre" if direction == "asc" else "-cliente__nombre",
            "cliente__apellido" if direction == "asc" else "-cliente__apellido",
        )

    elif sort == "estado":
        ventas = ventas.order_by("estado" if direction == "asc" else "-estado")

    elif sort == "total":
        ventas = ventas.order_by(
            "total_orden" if direction == "asc" else "-total_orden"
        )

    else:
        sort = "fecha"
        ventas = ventas.order_by("fecha" if direction == "asc" else "-fecha")

    return render(
        request,
        "ventas/lista_ventas.html",
        {
            "ventas": ventas,
            "q": q,
            "estado": estado,
            "sort": sort,
            "direction": direction,
            "current_sort": sort,
            "current_dir": direction,
        },
    )


@require_POST
def toggle_estado_venta(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.prefetch_related("detalles__producto", "detalles__devoluciones"),
        id=venta_id
    )
    nuevo_estado = "anulada" if venta.estado == "activa" else "activa"

    try:
        with transaction.atomic():
            for detalle in venta.detalles.all():
                if not detalle.producto:
                    continue
                cantidad_restante = detalle.cantidad_disponible_devolver
                if cantidad_restante <= 0:
                    continue

                delta = cantidad_restante if nuevo_estado == "anulada" else -cantidad_restante
                _ajustar_stock_producto(detalle.producto, delta)

            venta.estado = nuevo_estado
            venta.save()
    except ValidationError as exc:
        messages.error(request, str(exc))
        return redirect(reverse("ventas:lista") + "?estado=activa")

    messages.success(
        request, f"Estado actualizado correctamente: {venta.estado.upper()}"
    )
    return redirect(reverse("ventas:lista") + "?estado=activa")


def es_ajax(request):
    return request.headers.get("x-requested-with") == "XMLHttpRequest"


def render_crear_venta(request, form, productos_stock, servicios, personal, status=200):
    """
    - Si es AJAX: retorna HTML parcial para meterlo dentro del modal.
    - Si no es AJAX: retorna la página completa normal.
    """
    if es_ajax(request):
        html = render_to_string(
            "ventas/partials/crear_venta_form.html",
            {
                "form": form,
                "productos_stock": productos_stock,
                "servicios": servicios,
                "personal": personal,
            },
            request=request,
        )
        return JsonResponse({"success": False, "html": html}, status=status)

    return render(
        request,
        "ventas/crear_venta.html",
        {
            "form": form,
            "productos_stock": productos_stock,
            "servicios": servicios,
            "personal": personal,
        },
        status=status,
    )


@transaction.atomic
def crear_venta(request):
    productos = Producto.objects.filter(activo=True).select_related("stock")
    productos_stock = [
        {"producto": p, "stock": p.stock_actual, "activo": p.activo}
        for p in productos
    ]

    servicios = Servicio.objects.all()
    personal = Personal.objects.filter(rol="Colaborador", activo=True).order_by(
        "nombres", "apellidos"
    )
    if request.method == "POST":
        form = VentaForm(request.POST)
        items_json = request.POST.get("items")

        # 1) Validar items_json
        if not items_json:
            messages.error(
                request,
                "No se recibieron items. Agrega al menos 1 producto o servicio.",
            )
            return render(
                request,
                "ventas/crear_venta.html",
                {
                    "form": form,
                    "productos_stock": productos_stock,
                    "servicios": servicios,
                    "personal": personal,
                },
            )

        # 2) Convertir JSON a lista
        try:
            items = json.loads(items_json)
        except json.JSONDecodeError:
            messages.error(request, "El JSON de items llegó dañado. Revisa ventas.js.")
            return render(
                request,
                "ventas/crear_venta.html",
                {
                    "form": form,
                    "productos_stock": productos_stock,
                    "servicios": servicios,
                    "personal": personal,
                },
            )

        # 3) Debe haber items
        if not items:
            messages.error(
                request, "Agrega al menos 1 producto o servicio antes de guardar."
            )
            return render(
                request,
                "ventas/crear_venta.html",
                {
                    "form": form,
                    "productos_stock": productos_stock,
                    "servicios": servicios,
                    "personal": personal,
                },
            )

        # 4) Validar formulario
        if not form.is_valid():
            messages.error(request, "Formulario inválido. Revisa los campos.")
            return render(
                request,
                "ventas/crear_venta.html",
                {
                    "form": form,
                    "productos_stock": productos_stock,
                    "servicios": servicios,
                    "personal": personal,
                },
            )

        # ===============================
        # CREAR VENTA
        # ===============================
        venta = form.save(commit=False)
        venta.codigo_colaborador = "PENDIENTE"
        venta.nombre_colaborador = "Pendiente"

        # ✅ codigo_producto es obligatorio en tu modelo Venta
        # lo llenamos con un resumen simple
        it0 = items[0]
        if it0.get("tipo") == "producto":
            base = it0.get("id", "")
        else:
            base = it0.get("id_servicio", "")
        venta.codigo_producto = (
            str(base) if len(items) == 1 else f"{base} (+{len(items) - 1})"
        )

        venta.save()

        # ===============================
        # GUARDAR DETALLES
        # ===============================
        try:
            for item in items:
                if item.get("tipo") == "producto":
                    codigo = item.get("id")
                    if not codigo:
                        continue

                    producto = Producto.objects.select_for_update().get(codigo=codigo)
                    if not producto.activo:
                        raise ValidationError(
                            f"El producto {producto.nombre} no esta activo para ventas."
                        )

                    stock_real = producto.stock_actual

                    cantidad = int(item["cantidad"])
                    if cantidad > stock_real:
                        raise ValidationError(f"Stock insuficiente para {producto.nombre}")

                    DetalleVenta.objects.create(
                        venta=venta,
                        producto=producto,
                        precio_unitario=producto.precio,
                        cantidad=cantidad,
                        subtotal=producto.precio * cantidad,
                    )
                    _ajustar_stock_producto(producto, -cantidad)

                elif item.get("tipo") == "servicio":
                    servicio = Servicio.objects.get(id_servicio=item["id_servicio"])
                    colaborador = Personal.objects.get(id=item["id_personal"])

                    DetalleVenta.objects.create(
                        venta=venta,
                        servicio=servicio,
                        colaborador_servicio=colaborador,
                        precio_unitario=servicio.precio,
                        cantidad=int(item.get("cantidad", 1)),
                        subtotal=servicio.precio * int(item.get("cantidad", 1)),
                    )
        except ValidationError as exc:
            transaction.set_rollback(True)
            messages.error(request, str(exc))
            return render(
                request,
                "ventas/crear_venta.html",
                {
                    "form": form,
                    "productos_stock": productos_stock,
                    "servicios": servicios,
                    "personal": personal,
                },
            )

        if es_ajax(request):
            return JsonResponse({"success": True})

        messages.success(request, "Venta registrada correctamente.")
        return redirect("ventas:lista")

    # GET
    form = VentaForm()
    if es_ajax(request):
        html = render_to_string(
            "ventas/partials/crear_venta_form.html",
            {
                "form": form,
                "productos_stock": productos_stock,
                "servicios": servicios,
                "personal": personal,
            },
            request=request,
        )
        return JsonResponse({"success": True, "html": html})

    return render(
        request,
        "ventas/crear_venta.html",
        {
            "form": form,
            "productos_stock": productos_stock,
            "servicios": servicios,
            "personal": personal,
        },
    )


def editar_venta_modal(request, pk):
    venta = get_object_or_404(
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__producto", "detalles__devoluciones", "detalles__servicio"),
        pk=pk,
    )

    detalles_productos = venta.detalles.filter(producto__isnull=False).select_related("producto")
    detalles_servicios = venta.detalles.filter(servicio__isnull=False).select_related("servicio")
    servicios = Servicio.objects.all()
    personal = Personal.objects.filter(rol="Colaborador", activo=True).order_by("nombres", "apellidos")

    # Stock disponible para cada producto (desde Inventario.Stock)
    # Stock actual + lo que ya tiene esta venta (porque se va a editar)
    stock_por_producto = {}
    for det in detalles_productos:
        stock_actual = (
            Stock.objects.filter(producto=det.producto).values_list("cantidad_actual", flat=True).first()
            or 0
        )
        stock_por_producto[det.id] = stock_actual + det.cantidad_disponible_devolver

    if request.method == "POST":
        errores = []
        with transaction.atomic():
            # 1) Restaurar stock de los productos actuales
            for det in detalles_productos:
                _ajustar_stock_producto(det.producto, det.cantidad_disponible_devolver)

            # 2) Guardar cambios de productos
            for det in detalles_productos:
                cant_str = request.POST.get(f"prod_cant_{det.id}")
                nuevo_codigo = request.POST.get(f"prod_codigo_{det.id}")
                if cant_str is None:
                    continue

                nueva_cantidad = int(cant_str)

                # Cambio de producto si seleccionó uno diferente
                if nuevo_codigo and nuevo_codigo != det.producto.codigo:
                    det.producto = Producto.objects.get(codigo=nuevo_codigo)

                # El precio se toma siempre del producto (no editable)
                nuevo_precio = det.producto.precio

                # Validar stock desde Inventario (ya restauramos el stock al inicio)
                disponible = (
                    Stock.objects.select_for_update()
                    .filter(producto=det.producto)
                    .values_list("cantidad_actual", flat=True)
                    .first()
                    or 0
                )
                if nueva_cantidad > disponible:
                    nombre = det.producto.nombre
                    return JsonResponse(
                        {"ok": False, "error": f"Stock insuficiente para '{nombre}'. Disponible: {disponible}."},
                        status=400,
                    )

                det.cantidad = nueva_cantidad
                det.precio_unitario = nuevo_precio
                det.subtotal = det.cantidad * det.precio_unitario
                det.save()

                # 3) Descontar el nuevo stock
                _ajustar_stock_producto(det.producto, -nueva_cantidad)

            # 4) Guardar cambios de servicios
            for det in detalles_servicios:
                cant_str = request.POST.get(f"serv_cant_{det.id}")
                serv_id = request.POST.get(f"serv_servicio_{det.id}")
                pers_id = request.POST.get(f"serv_personal_{det.id}")

                if serv_id:
                    det.servicio = Servicio.objects.get(pk=serv_id)
                if pers_id:
                    det.colaborador_servicio = Personal.objects.get(pk=pers_id)
                if cant_str:
                    det.cantidad = int(cant_str)

                det.precio_unitario = det.servicio.precio
                det.subtotal = det.cantidad * det.precio_unitario
                det.save()

        return JsonResponse({"ok": True})

    # GET — anotar cada detalle con su stock para usarlo directo en el template
    for det in detalles_productos:
        det.stock_disponible = stock_por_producto.get(det.id, 0)

    # Lista completa de productos con su stock (para el select de cambio de producto)
    productos_relacionados_ids = detalles_productos.values_list("producto_id", flat=True)
    todos_productos = Producto.objects.filter(
        Q(activo=True) | Q(codigo__in=productos_relacionados_ids)
    ).distinct()
    todos_stock = []
    for p in todos_productos:
        stock_actual = (
            Stock.objects.filter(producto=p).values_list("cantidad_actual", flat=True).first() or 0
        )
        # Si este producto ya está en la venta, sumar lo que tiene (para permitir editar sin bloquearse)
        det_actual = detalles_productos.filter(producto=p).first()
        ya_tiene = det_actual.cantidad_disponible_devolver if det_actual else 0
        s = stock_actual + ya_tiene
        todos_stock.append({"producto": p, "stock": s, "activo": p.activo})

    ctx = {
        "venta": venta,
        "detalles_productos": detalles_productos,
        "detalles_servicios": detalles_servicios,
        "servicios": servicios,
        "personal": personal,
        "todos_stock": todos_stock,
    }

    if es_ajax(request):
        html = render_to_string("ventas/form_editar_venta.html", ctx, request=request)
        return JsonResponse({"success": True, "html": html})

    return render(request, "ventas/form_editar_venta.html", ctx)

def detalle_venta_json(request, pk):
    venta = get_object_or_404(
        Venta.objects.select_related("cliente").prefetch_related(
            "detalles__producto",
            "detalles__servicio",
            "detalles__colaborador_servicio",
        ),
        pk=pk,
    )

    detalles = []
    for d in venta.detalles.all():
        nombre_item = ""
        tipo = ""

        if d.producto:
            nombre_item = d.producto.nombre
            tipo = "Producto"
        elif d.servicio:
            nombre_item = d.servicio.nombre
            tipo = "Servicio"

        detalles.append(
            {
                "tipo": tipo,
                "nombre": nombre_item,
                "cantidad": d.cantidad,
                "precio_unitario": str(d.precio_unitario),
                "subtotal": str(d.subtotal),
                "colaborador": str(d.colaborador_servicio)
                if d.colaborador_servicio
                else "",
            }
        )

    data = {
        "id": venta.id,
        "codigo_venta": venta.codigo_venta,
        "cliente": str(venta.cliente),
        "fecha": venta.fecha.strftime("%d/%m/%Y %H:%M") if venta.fecha else "",
        "estado": venta.estado,
        "total": str(venta.total),
        "detalles": detalles,
    }

    return JsonResponse(data)


@transaction.atomic
def anular_venta(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.prefetch_related("detalles__producto", "detalles__devoluciones"),
        id=venta_id
    )

    if venta.estado == "anulada":
        return redirect(reverse("ventas:lista") + "?estado=activa")

    if request.method == "POST":
        for detalle in venta.detalles.all():
            if detalle.producto:
                cantidad_restante = detalle.cantidad_disponible_devolver
                if cantidad_restante > 0:
                    _ajustar_stock_producto(detalle.producto, cantidad_restante)
        venta.estado = "anulada"
        venta.save()

    return redirect(reverse("ventas:lista") + "?estado=activa")


COLUMNAS_REPORTE_VENTAS = {
    "codigo_venta": "Código venta",
    "cliente": "Cliente",
    "productos_servicios": "Productos / Servicios",
    "fecha": "Fecha",
    "total": "Total",
    "estado": "Estado",
}


def construir_queryset_reporte_ventas(request):
    fecha_inicio = request.GET.get("fecha_inicio", "").strip()
    fecha_fin = request.GET.get("fecha_fin", "").strip()
    fecha_inicio_comp = request.GET.get("fecha_inicio_comp", "").strip()
    fecha_fin_comp = request.GET.get("fecha_fin_comp", "").strip()

    incluir_total = request.GET.get("incluir_total") == "1"
    comparativo = request.GET.get("comparativo") == "1"
    incluir_grafica = (request.GET.get("incluir_grafica") == "1") and comparativo
    tipo_grafica = request.GET.get("tipo_grafica", "bar").strip()
    if tipo_grafica not in {"bar", "line", "pie"}:
        tipo_grafica = "bar"
    columnas = request.GET.getlist("columnas")

    ventas = (
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__producto", "detalles__servicio")
        .annotate(
            total_orden=Coalesce(
                Sum("detalles__subtotal"),
                Value(0),
                output_field=DecimalField(max_digits=12, decimal_places=2),
            )
        )
        .order_by("fecha")
    )

    if fecha_inicio and fecha_fin:
        ventas = ventas.filter(fecha__date__range=[fecha_inicio, fecha_fin])

    ventas_comp = Venta.objects.none()
    if comparativo and fecha_inicio_comp and fecha_fin_comp:
        ventas_comp = (
            Venta.objects.select_related("cliente")
            .prefetch_related("detalles__producto", "detalles__servicio")
            .annotate(
                total_orden=Coalesce(
                    Sum("detalles__subtotal"),
                    Value(0),
                    output_field=DecimalField(max_digits=12, decimal_places=2),
                )
            )
            .filter(fecha__date__range=[fecha_inicio_comp, fecha_fin_comp])
            .order_by("fecha")
        )

    total_principal = ventas.aggregate(
        total=Coalesce(
            Sum("detalles__subtotal"),
            Value(0),
            output_field=DecimalField(max_digits=12, decimal_places=2),
        )
    )["total"]

    total_comparativo = 0
    if comparativo and fecha_inicio_comp and fecha_fin_comp:
        total_comparativo = ventas_comp.aggregate(
            total=Coalesce(
                Sum("detalles__subtotal"),
                Value(0),
                output_field=DecimalField(max_digits=12, decimal_places=2),
            )
        )["total"]

    # Las gráficas solo se generan cuando hay comparativo activo
    # (para mantener el mismo comportamiento solicitado: gráfica únicamente en modo comparativo).
    incluir_grafica = bool(incluir_grafica and fecha_inicio_comp and fecha_fin_comp)

    grafica_principal_labels = []
    grafica_principal_data = []
    if incluir_grafica:
        grafica_principal = (
            ventas.annotate(dia=TruncDate("fecha"))
            .values("dia")
            .annotate(
                total=Coalesce(
                    Sum("detalles__subtotal"),
                    Value(0),
                    output_field=DecimalField(max_digits=12, decimal_places=2),
                )
            )
            .order_by("dia")
        )
        grafica_principal_labels = [
            item["dia"].strftime("%d/%m/%Y")
            for item in grafica_principal
            if item["dia"]
        ]
        grafica_principal_data = [float(item["total"]) for item in grafica_principal]

    grafica_comp_labels = []
    grafica_comp_data = []
    if incluir_grafica and comparativo and fecha_inicio_comp and fecha_fin_comp:
        grafica_comp = (
            ventas_comp.annotate(dia=TruncDate("fecha"))
            .values("dia")
            .annotate(
                total=Coalesce(
                    Sum("detalles__subtotal"),
                    Value(0),
                    output_field=DecimalField(max_digits=12, decimal_places=2),
                )
            )
            .order_by("dia")
        )
        grafica_comp_labels = [
            item["dia"].strftime("%d/%m/%Y") for item in grafica_comp if item["dia"]
        ]
        grafica_comp_data = [float(item["total"]) for item in grafica_comp]

    return {
        "ventas": ventas,
        "ventas_comp": ventas_comp,
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "fecha_inicio_comp": fecha_inicio_comp,
        "fecha_fin_comp": fecha_fin_comp,
        "incluir_total": incluir_total,
        "incluir_grafica": incluir_grafica,
        "comparativo": comparativo,
        "tipo_grafica": tipo_grafica,
        "columnas": columnas,
        "total_principal": total_principal,
        "total_comparativo": total_comparativo,
        "grafica_principal_labels": grafica_principal_labels,
        "grafica_principal_data": grafica_principal_data,
        "grafica_comp_labels": grafica_comp_labels,
        "grafica_comp_data": grafica_comp_data,
    }



def obtener_valor_columna_venta(venta, columna):
    if columna == "codigo_venta":
        return venta.codigo_venta
    if columna == "cliente":
        return str(venta.cliente)
    if columna == "productos_servicios":
        items = []
        for d in venta.detalles.all():
            nombre = ""
            if d.producto:
                nombre = d.producto.nombre
                if not d.producto.activo:
                    nombre += " (Inactivo)"
            elif d.servicio:
                nombre = d.servicio.nombre
            if nombre:
                items.append(f"{nombre} x{d.cantidad}")
        return ", ".join(items) if items else "Sin detalles"
    if columna == "fecha":
        return venta.fecha.strftime("%d/%m/%Y %H:%M") if venta.fecha else ""
    if columna == "total":
        return f"{venta.total:.2f}"
    if columna == "estado":
        return venta.estado.title()
    return ""


def vista_previa_reporte_ventas(request):
    data = construir_queryset_reporte_ventas(request)
    html = render_to_string(
        "ventas/partials/reporte_ventas_preview.html",
        {
            **data,
            "columnas_map": COLUMNAS_REPORTE_VENTAS,
        },
        request=request,
    )

    return JsonResponse(
        {
            "html": html,
            "grafica_principal_labels": data["grafica_principal_labels"],
            "grafica_principal_data": data["grafica_principal_data"],
            "grafica_comp_labels": data["grafica_comp_labels"],
            "grafica_comp_data": data["grafica_comp_data"],
            "incluir_grafica": data["incluir_grafica"],
            "comparativo": data["comparativo"],
            "tipo_grafica": data["tipo_grafica"],
            "fecha_inicio": data["fecha_inicio"],
            "fecha_fin": data["fecha_fin"],
            "fecha_inicio_comp": data["fecha_inicio_comp"],
            "fecha_fin_comp": data["fecha_fin_comp"],
        }
    )


def exportar_reporte_ventas(request):
    data = construir_queryset_reporte_ventas(request)
    formato = request.GET.get("formato", "pdf").strip().lower()
    columnas = data["columnas"] or [
        "codigo_venta",
        "cliente",
        "fecha",
        "total",
        "estado",
    ]

    if formato == "excel":
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte Ventas"

        def apply_style(cell, font=None, fill=None, border=None, alignment=None, number_format=None):
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            if border:
                cell.border = border
            if alignment:
                cell.alignment = alignment
            if number_format:
                cell.number_format = number_format

        def style_range(start_row, end_row, start_col, end_col, **styles):
            for r in range(start_row, end_row + 1):
                for c in range(start_col, end_col + 1):
                    apply_style(ws.cell(row=r, column=c), **styles)

        # Configuración visual de hoja (mismo look que Compras)
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = 0.35
        ws.page_margins.right = 0.35
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4

        color_dark = "3B261A"
        color_brown = "A67C52"
        color_soft = "F7F0E8"
        color_soft_2 = "FBF7F2"
        color_line = "D8C3B5"
        color_white = "FFFFFF"
        color_text = "2F241D"

        title_font = Font(bold=True, size=16, color=color_text)
        section_font = Font(bold=True, size=11, color=color_white)
        label_font = Font(bold=True, size=10, color=color_dark)
        value_font = Font(size=10, color=color_text)
        value_bold_font = Font(bold=True, size=10, color=color_text)
        table_header_font = Font(bold=True, size=10, color=color_white)
        footer_font = Font(italic=True, size=9, color=color_brown)

        dark_fill = PatternFill("solid", fgColor=color_dark)
        brown_fill = PatternFill("solid", fgColor=color_brown)
        soft_fill = PatternFill("solid", fgColor=color_soft)
        soft_fill_2 = PatternFill("solid", fgColor=color_soft_2)

        thin_side = Side(style="thin", color=color_line)
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        right = Alignment(horizontal="right", vertical="center", wrap_text=True)

        ws.column_dimensions["A"].width = 3
        width_map = {
            "codigo_venta": 16,
            "cliente": 26,
            "productos_servicios": 48,
            "fecha": 20,
            "total": 16,
            "estado": 14,
        }
        for i, key in enumerate(columnas, start=2):
            ws.column_dimensions[get_column_letter(i)].width = width_map.get(key, 20)

        last_col = len(columnas) + 1  # B..

        # Logo
        logo_path = _get_brand_logo_path()
        if logo_path:
            try:
                img = XLImage(logo_path)
                img.width = 150
                img.height = 58
                ws.add_image(img, "B1")
            except Exception:
                pass

        # Título
        style_range(1, 2, 2, last_col, border=thin_border, alignment=center)
        ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=last_col)
        ws["B1"] = "Reporte de Ventas"
        apply_style(ws["B1"], font=title_font, alignment=center)

        # Datos generales
        row = 4
        style_range(row, row, 2, last_col, fill=dark_fill, border=thin_border, alignment=left)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
        ws.cell(row=row, column=2, value="DATOS GENERALES")
        apply_style(ws.cell(row=row, column=2), font=section_font, alignment=left)

        row += 1
        info_rows = [
            ("Rango principal", f"{data['fecha_inicio']} - {data['fecha_fin']}"),
        ]
        if data["comparativo"]:
            info_rows.append(("Rango comparativo", f"{data['fecha_inicio_comp']} - {data['fecha_fin_comp']}"))
        if data["incluir_total"]:
            info_rows.append(("Total rango principal", _fmt_money_excel(data["total_principal"])))
            if data["comparativo"]:
                info_rows.append(("Total rango comparativo", _fmt_money_excel(data["total_comparativo"])))

        for label, value in info_rows:
            ws.cell(row=row, column=2, value=label)
            apply_style(ws.cell(row=row, column=2), font=label_font, fill=soft_fill, border=thin_border, alignment=left)

            style_range(row, row, 3, last_col, fill=soft_fill_2, border=thin_border, alignment=left)
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=last_col)
            ws.cell(row=row, column=3, value=value)

            if isinstance(value, (int, float)) and "Total" in label:
                apply_style(
                    ws.cell(row=row, column=3),
                    font=value_bold_font,
                    fill=soft_fill_2,
                    border=thin_border,
                    alignment=right,
                    number_format='$#,##0.00',
                )
            else:
                apply_style(ws.cell(row=row, column=3), font=value_font, fill=soft_fill_2, border=thin_border, alignment=left)

            row += 1

        # Sección detalle
        row += 1
        style_range(row, row, 2, last_col, fill=dark_fill, border=thin_border, alignment=left)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
        ws.cell(row=row, column=2, value="DETALLE DE VENTAS")
        apply_style(ws.cell(row=row, column=2), font=section_font, alignment=left)

        row += 1
        for idx, key in enumerate(columnas, start=2):
            cell = ws.cell(row=row, column=idx, value=COLUMNAS_REPORTE_VENTAS[key])
            apply_style(cell, font=table_header_font, fill=brown_fill, border=thin_border, alignment=center)

        row += 1
        for venta in data["ventas"]:
            for idx, key in enumerate(columnas, start=2):
                if key == "total":
                    value = _fmt_money_excel(getattr(venta, "total", 0))
                else:
                    value = obtener_valor_columna_venta(venta, key)
                ws.cell(row=row, column=idx, value=value)

            row_fill = soft_fill if row % 2 == 1 else soft_fill_2
            for idx, key in enumerate(columnas, start=2):
                cell = ws.cell(row=row, column=idx)
                align = left if key in ("cliente", "productos_servicios") else center
                num_fmt = '$#,##0.00' if key == "total" else None
                apply_style(cell, fill=row_fill, border=thin_border, alignment=align, number_format=num_fmt)
                apply_style(cell, font=value_bold_font if key == "codigo_venta" else value_font)
            row += 1

        if not data["ventas"]:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
            ws.cell(row=row, column=2, value="No hay ventas para este rango.")
            style_range(row, row, 2, last_col, fill=soft_fill_2, border=thin_border, alignment=center)
            apply_style(ws.cell(row=row, column=2), font=value_font, alignment=center)
            row += 1

        # Footer
        row += 1
        style_range(row, row, 2, last_col, alignment=center)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
        ws.cell(row=row, column=2, value="Generado por MonaApp / Monakeratina")
        apply_style(ws.cell(row=row, column=2), font=footer_font, alignment=center)

        ws.print_area = f"B1:{get_column_letter(last_col)}{row}"

        # Mantener hojas de gráficas existentes (si aplica)
        if data["incluir_grafica"] and data["grafica_principal_labels"]:
            ws_chart = wb.create_sheet("Grafica")
            ws_chart.append(["Fecha", "Total"])
            for lbl, val in zip(data["grafica_principal_labels"], data["grafica_principal_data"]):
                ws_chart.append([lbl, val])

            tipo = data.get("tipo_grafica", "bar")
            chart = LineChart() if tipo == "line" else BarChart()
            chart.title = "Ventas - rango principal"
            chart.y_axis.title = "Total ($)"
            chart.x_axis.title = "Fecha"
            chart.style = 10

            data_ref = Reference(ws_chart, min_col=2, min_row=1, max_row=len(data["grafica_principal_labels"]) + 1)
            cats_ref = Reference(ws_chart, min_col=1, min_row=2, max_row=len(data["grafica_principal_labels"]) + 1)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            ws_chart.add_chart(chart, "D2")

            if data["comparativo"] and data["grafica_comp_labels"]:
                ws_chart2 = wb.create_sheet("Grafica comparativo")
                ws_chart2.append(["Fecha", "Total"])
                for lbl, val in zip(data["grafica_comp_labels"], data["grafica_comp_data"]):
                    ws_chart2.append([lbl, val])

                chart2 = LineChart() if tipo == "line" else BarChart()
                chart2.title = "Ventas - rango comparativo"
                chart2.y_axis.title = "Total ($)"
                chart2.style = 10
                data_ref2 = Reference(ws_chart2, min_col=2, min_row=1, max_row=len(data["grafica_comp_labels"]) + 1)
                cats_ref2 = Reference(ws_chart2, min_col=1, min_row=2, max_row=len(data["grafica_comp_labels"]) + 1)
                chart2.add_data(data_ref2, titles_from_data=True)
                chart2.set_categories(cats_ref2)
                ws_chart2.add_chart(chart2, "D2")

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="reporte_ventas.xlsx"'
        return response

    if formato == "_excel_legacy":
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte Ventas"

        encabezados = [COLUMNAS_REPORTE_VENTAS[c] for c in columnas]
        ws.append(encabezados)

        for cell in ws[1]:
            cell.font = Font(bold=True)

        for venta in data["ventas"]:
            ws.append([obtener_valor_columna_venta(venta, c) for c in columnas])

        if data["incluir_total"]:
            ws.append([])
            ws.append(["Total rango principal", f"{data['total_principal']:.2f}"])
            if data["comparativo"]:
                ws.append(
                    ["Total rango comparativo", f"{data['total_comparativo']:.2f}"]
                )

        if data["incluir_grafica"] and data["grafica_principal_labels"]:
            ws_chart = wb.create_sheet("Gráfica")
            ws_chart.append(["Fecha", "Total"])
            for lbl, val in zip(data["grafica_principal_labels"], data["grafica_principal_data"]):
                ws_chart.append([lbl, val])

            tipo = data.get("tipo_grafica", "bar")
            chart = LineChart() if tipo == "line" else BarChart()
            chart.title = "Ventas — rango principal"
            chart.y_axis.title = "Total ($)"
            chart.x_axis.title = "Fecha"
            chart.style = 10

            data_ref = Reference(ws_chart, min_col=2, min_row=1, max_row=len(data["grafica_principal_labels"]) + 1)
            cats_ref = Reference(ws_chart, min_col=1, min_row=2, max_row=len(data["grafica_principal_labels"]) + 1)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            chart.shape = 4
            ws_chart.add_chart(chart, "D2")

            if data["comparativo"] and data["grafica_comp_labels"]:
                ws_chart2 = wb.create_sheet("Gráfica comparativo")
                ws_chart2.append(["Fecha", "Total"])
                for lbl, val in zip(data["grafica_comp_labels"], data["grafica_comp_data"]):
                    ws_chart2.append([lbl, val])

                chart2 = LineChart() if tipo == "line" else BarChart()
                chart2.title = "Ventas — rango comparativo"
                chart2.y_axis.title = "Total ($)"
                chart2.style = 10
                data_ref2 = Reference(ws_chart2, min_col=2, min_row=1, max_row=len(data["grafica_comp_labels"]) + 1)
                cats_ref2 = Reference(ws_chart2, min_col=1, min_row=2, max_row=len(data["grafica_comp_labels"]) + 1)
                chart2.add_data(data_ref2, titles_from_data=True)
                chart2.set_categories(cats_ref2)
                ws_chart2.add_chart(chart2, "D2")

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="reporte_ventas.xlsx"'
        return response

    # =========================
    # PDF con estilo tipo Compras
    # =========================
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()
    elements = []

    color_dark = colors.HexColor("#3B261A")
    color_brown = colors.HexColor("#A67C52")
    color_soft = colors.HexColor("#F7F0E8")
    color_soft_2 = colors.HexColor("#FBF7F2")
    color_line = colors.HexColor("#D8C3B5")
    color_text = colors.HexColor("#2F241D")

    title_style = ParagraphStyle(
        "ReporteVentasTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.black,
        alignment=1,
    )

    label_style = ParagraphStyle(
        "ReporteVentasLabel",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=9,
        textColor=color_text,
    )

    value_style = ParagraphStyle(
        "ReporteVentasValue",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        leading=11,
        textColor=color_text,
    )

    cell_style = ParagraphStyle(
        "ReporteVentasCell",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=color_text,
    )

    # Header con logo + título
    logo_path = _get_brand_logo_path()
    logo_flowable = ""
    if logo_path:
        try:
            logo_flowable = RLImage(logo_path, width=34 * mm, height=16 * mm)
        except Exception:
            logo_flowable = ""

    header_table = Table(
        [[logo_flowable, Paragraph("Reporte de Ventas", title_style), ""]],
        colWidths=[42 * mm, doc.width - (42 * mm + 28 * mm), 28 * mm],
    )
    header_table.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (0, 0), "LEFT"),
                ("ALIGN", (1, 0), (1, 0), "CENTER"),
                ("LINEBELOW", (0, 0), (-1, 0), 1.2, colors.black),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    elements.append(header_table)
    elements.append(Spacer(1, 8))

    # Datos generales
    info_head = Table([[Paragraph("DATOS GENERALES", label_style)]], colWidths=[doc.width])
    info_head.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), color_dark),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("BOX", (0, 0), (-1, -1), 0.8, color_line),
            ]
        )
    )
    elements.append(info_head)

    info_rows = [
        [Paragraph("Rango principal", label_style), Paragraph(f"{data['fecha_inicio']} - {data['fecha_fin']}", value_style)],
    ]
    if data["comparativo"]:
        info_rows.append([Paragraph("Rango comparativo", label_style), Paragraph(f"{data['fecha_inicio_comp']} - {data['fecha_fin_comp']}", value_style)])
    if data["incluir_total"]:
        info_rows.append([Paragraph("Total principal", label_style), Paragraph(_fmt_money_text(data["total_principal"]), value_style)])
        if data["comparativo"]:
            info_rows.append([Paragraph("Total comparativo", label_style), Paragraph(_fmt_money_text(data["total_comparativo"]), value_style)])

    info_table = Table(info_rows, colWidths=[45 * mm, doc.width - 45 * mm])
    info_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.6, color_line),
                ("BACKGROUND", (0, 0), (-1, -1), color_soft_2),
                ("BACKGROUND", (0, 0), (0, -1), color_soft),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    elements.append(info_table)
    elements.append(Spacer(1, 10))

    def _col_widths(keys):
        base_map = {
            "codigo_venta": 28 * mm,
            "cliente": 45 * mm,
            "productos_servicios": 95 * mm,
            "fecha": 36 * mm,
            "total": 26 * mm,
            "estado": 22 * mm,
        }
        widths = [base_map.get(k, 32 * mm) for k in keys]
        total_w = sum(widths)
        if total_w <= doc.width:
            return widths
        factor = float(doc.width) / float(total_w)
        return [w * factor for w in widths]

    header_labels = [COLUMNAS_REPORTE_VENTAS[c] for c in columnas]
    tabla_data = [[Paragraph(str(h), label_style) for h in header_labels]]

    for venta in data["ventas"]:
        row = []
        for c in columnas:
            if c == "total":
                row.append(Paragraph(_fmt_money_text(getattr(venta, "total", 0)), cell_style))
            else:
                row.append(Paragraph(str(obtener_valor_columna_venta(venta, c)), cell_style))
        tabla_data.append(row)

    if not data["ventas"]:
        tabla_data.append([Paragraph("No hay ventas para este rango.", cell_style)] + [""] * (len(columnas) - 1))

    tabla = Table(tabla_data, repeatRows=1, colWidths=_col_widths(columnas))
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), color_brown),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, color_line),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [color_soft_2, color_soft]),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    elements.append(tabla)

    # Gráficas (mismo contenido, colores más premium)
    if data["incluir_grafica"] and data["grafica_principal_labels"]:
        elements.append(Spacer(1, 14))
        elements.append(Paragraph("Gráfica de ventas - rango principal", styles["Heading3"]))
        elements.append(Spacer(1, 6))

        tipo = data.get("tipo_grafica", "bar")
        labels = data["grafica_principal_labels"]
        valores = data["grafica_principal_data"]
        img_buf = _build_plotly_chart_image(
            labels,
            valores,
            tipo=tipo,
            title="Ventas - rango principal",
            color="#A67C52",
        )
        if img_buf:
            elements.append(RLImage(img_buf, width=560, height=220))
        else:
            elements.append(Paragraph("No fue posible generar la gráfica del rango principal.", styles["Italic"]))

        if data["comparativo"] and data["grafica_comp_labels"]:
            elements.append(Spacer(1, 14))
            elements.append(Paragraph("Gráfica de ventas - rango comparativo", styles["Heading3"]))
            elements.append(Spacer(1, 6))

            labels2 = data["grafica_comp_labels"]
            valores2 = data["grafica_comp_data"]
            img_buf2 = _build_plotly_chart_image(
                labels2,
                valores2,
                tipo=tipo,
                title="Ventas - rango comparativo",
                color="#3B261A",
            )
            if img_buf2:
                elements.append(RLImage(img_buf2, width=560, height=220))
            else:
                elements.append(Paragraph("No fue posible generar la gráfica comparativa.", styles["Italic"]))

    doc.build(elements, onFirstPage=_draw_reporte_watermark, onLaterPages=_draw_reporte_watermark)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="reporte_ventas.pdf"'
    return response

    # LEGACY (sin uso, se deja por referencia)
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter))
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Reporte de Ventas", styles["Title"]))
    elements.append(Spacer(1, 12))
    elements.append(
        Paragraph(
            f"Rango principal: {data['fecha_inicio']} a {data['fecha_fin']}",
            styles["Normal"],
        )
    )
    if data["comparativo"]:
        elements.append(
            Paragraph(
                f"Rango comparativo: {data['fecha_inicio_comp']} a {data['fecha_fin_comp']}",
                styles["Normal"],
            )
        )
    elements.append(Spacer(1, 12))

    tabla_data = [[COLUMNAS_REPORTE_VENTAS[c] for c in columnas]]
    for venta in data["ventas"]:
        tabla_data.append([obtener_valor_columna_venta(venta, c) for c in columnas])

    tabla = Table(tabla_data, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#8d604a")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.whitesmoke, colors.lightgrey],
                ),
            ]
        )
    )
    elements.append(tabla)

    if data["incluir_total"]:
        elements.append(Spacer(1, 12))
        elements.append(
            Paragraph(
                f"Total rango principal: {data['total_principal']:.2f}",
                styles["Heading3"],
            )
        )
        if data["comparativo"]:
            elements.append(
                Paragraph(
                    f"Total rango comparativo: {data['total_comparativo']:.2f}",
                    styles["Heading3"],
                )
            )

    if data["incluir_grafica"] and data["grafica_principal_labels"]:
        elements.append(Spacer(1, 16))
        elements.append(Paragraph("Gráfica de ventas — rango principal", styles["Heading3"]))
        elements.append(Spacer(1, 6))

        tipo = data.get("tipo_grafica", "bar")
        labels = data["grafica_principal_labels"]
        valores = data["grafica_principal_data"]
        img_buf = _build_plotly_chart_image(
            labels,
            valores,
            tipo=tipo,
            title="Ventas — rango principal",
            color="#8d604a",
        )
        if img_buf:
            elements.append(RLImage(img_buf, width=560, height=220))
        else:
            elements.append(Paragraph("No fue posible generar la gráfica del rango principal.", styles["Italic"]))

        if data["comparativo"] and data["grafica_comp_labels"]:
            elements.append(Spacer(1, 16))
            elements.append(Paragraph("Gráfica de ventas — rango comparativo", styles["Heading3"]))
            elements.append(Spacer(1, 6))

            labels2 = data["grafica_comp_labels"]
            valores2 = data["grafica_comp_data"]
            img_buf2 = _build_plotly_chart_image(
                labels2,
                valores2,
                tipo=tipo,
                title="Ventas — rango comparativo",
                color="#4a7c8d",
            )
            if img_buf2:
                elements.append(RLImage(img_buf2, width=560, height=220))
            else:
                elements.append(Paragraph("No fue posible generar la gráfica comparativa.", styles["Italic"]))

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="reporte_ventas.pdf"'
    return response

# ============================================================
# DEVOLUCIÓN DE VENTAS
# ============================================================


def devolucion_venta_json(request, venta_id):
    """
    GET: Devuelve en JSON los ítems de la venta con cuánto se puede devolver.
    """
    venta = get_object_or_404(
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__producto", "detalles__servicio", "detalles__devoluciones"),
        id=venta_id
    )

    # Validaciones de negocio
    if venta.estado == "anulada":
        return JsonResponse({"error": "No se puede devolver una venta anulada."}, status=400)

    items = []
    for d in venta.detalles.all():
        ya_devuelto = d.cantidad_devuelta
        disponible = d.cantidad_disponible_devolver

        nombre = ""
        tipo = ""
        if d.producto:
            nombre = d.producto.nombre
            tipo = "producto"
        elif d.servicio:
            nombre = d.servicio.nombre
            tipo = "servicio"

        items.append({
            "detalle_id": d.id,
            "nombre": nombre,
            "tipo": tipo,
            "precio_unitario": float(d.precio_unitario),
            "cantidad_original": d.cantidad,
            "ya_devuelto": ya_devuelto,
            "disponible": disponible,
        })

    return JsonResponse({
        "venta_id": venta.id,
        "codigo_venta": venta.codigo_venta,
        "cliente": str(venta.cliente),
        "total_venta": float(venta.total),
        "total_ya_devuelto": float(venta.total_devuelto),
        "items": items,
    })


def editar_venta_modal(request, pk):
    venta = get_object_or_404(
        Venta.objects.select_related("cliente").prefetch_related(
            "detalles__producto",
            "detalles__devoluciones",
            "detalles__servicio",
        ),
        pk=pk,
    )

    detalles_productos = venta.detalles.filter(producto__isnull=False).select_related("producto")
    detalles_servicios = venta.detalles.filter(servicio__isnull=False).select_related("servicio")
    servicios = Servicio.objects.all()
    personal = Personal.objects.filter(rol="Colaborador", activo=True).order_by("nombres", "apellidos")

    stock_por_producto = {}
    for det in detalles_productos:
        stock_actual = (
            Stock.objects.filter(producto=det.producto).values_list("cantidad_actual", flat=True).first()
            or 0
        )
        stock_por_producto[det.id] = stock_actual + det.cantidad_disponible_devolver

    if request.method == "POST":
        try:
            with transaction.atomic():
                for det in detalles_productos:
                    _ajustar_stock_producto(det.producto, det.cantidad_disponible_devolver)

                for det in detalles_productos:
                    cant_str = request.POST.get(f"prod_cant_{det.id}")
                    nuevo_codigo = request.POST.get(f"prod_codigo_{det.id}")
                    if cant_str is None:
                        continue

                    nueva_cantidad = int(cant_str)
                    if nuevo_codigo and nuevo_codigo != det.producto.codigo:
                        det.producto = Producto.objects.get(codigo=nuevo_codigo)

                    disponible = (
                        Stock.objects.select_for_update()
                        .filter(producto=det.producto)
                        .values_list("cantidad_actual", flat=True)
                        .first()
                        or 0
                    )
                    if nueva_cantidad > disponible:
                        nombre = det.producto.nombre
                        return JsonResponse(
                            {"ok": False, "error": f"Stock insuficiente para '{nombre}'. Disponible: {disponible}."},
                            status=400,
                        )

                    det.cantidad = nueva_cantidad
                    det.precio_unitario = det.producto.precio
                    det.subtotal = det.cantidad * det.precio_unitario
                    det.save()

                    _ajustar_stock_producto(det.producto, -nueva_cantidad)

                for det in detalles_servicios:
                    cant_str = request.POST.get(f"serv_cant_{det.id}")
                    serv_id = request.POST.get(f"serv_servicio_{det.id}")
                    pers_id = request.POST.get(f"serv_personal_{det.id}")

                    if serv_id:
                        det.servicio = Servicio.objects.get(pk=serv_id)
                    if pers_id:
                        det.colaborador_servicio = Personal.objects.get(pk=pers_id)
                    if cant_str:
                        det.cantidad = int(cant_str)

                    det.precio_unitario = det.servicio.precio
                    det.subtotal = det.cantidad * det.precio_unitario
                    det.save()
        except ValidationError as exc:
            return JsonResponse({"ok": False, "error": str(exc)}, status=400)

        return JsonResponse({"ok": True})

    for det in detalles_productos:
        det.stock_disponible = stock_por_producto.get(det.id, 0)

    productos_relacionados_ids = detalles_productos.values_list("producto_id", flat=True)
    todos_productos = Producto.objects.filter(
        Q(activo=True) | Q(codigo__in=productos_relacionados_ids)
    ).distinct()
    todos_stock = []
    for p in todos_productos:
        stock_actual = (
            Stock.objects.filter(producto=p).values_list("cantidad_actual", flat=True).first() or 0
        )
        det_actual = detalles_productos.filter(producto=p).first()
        ya_tiene = det_actual.cantidad_disponible_devolver if det_actual else 0
        todos_stock.append({
            "producto": p,
            "stock": stock_actual + ya_tiene,
            "activo": p.activo,
        })

    ctx = {
        "venta": venta,
        "detalles_productos": detalles_productos,
        "detalles_servicios": detalles_servicios,
        "servicios": servicios,
        "personal": personal,
        "todos_stock": todos_stock,
    }

    if es_ajax(request):
        html = render_to_string("ventas/form_editar_venta.html", ctx, request=request)
        return JsonResponse({"success": True, "html": html})

    return render(request, "ventas/form_editar_venta.html", ctx)


@transaction.atomic
def anular_venta(request, venta_id):
    venta = get_object_or_404(
        Venta.objects.prefetch_related("detalles__producto", "detalles__devoluciones"),
        id=venta_id,
    )

    if venta.estado == "anulada":
        return redirect(reverse("ventas:lista") + "?estado=activa")

    if request.method == "POST":
        try:
            with transaction.atomic():
                for detalle in venta.detalles.all():
                    if detalle.producto:
                        cantidad_restante = detalle.cantidad_disponible_devolver
                        if cantidad_restante > 0:
                            _ajustar_stock_producto(detalle.producto, cantidad_restante)
                venta.estado = "anulada"
                venta.save()
        except ValidationError as exc:
            messages.error(request, str(exc))
            return redirect(reverse("ventas:lista") + "?estado=activa")

    return redirect(reverse("ventas:lista") + "?estado=activa")


@transaction.atomic
def registrar_devolucion(request, venta_id):
    """
    POST: Registra la devolución con validaciones completas.
    """
    if request.method != "POST":
        return JsonResponse({"error": "Método no permitido."}, status=405)

    venta = get_object_or_404(
        Venta.objects.select_related("cliente")
        .prefetch_related("detalles__devoluciones"),
        id=venta_id
    )

    # ── Validación 1: venta activa ──────────────────────────
    if venta.estado == "anulada":
        return JsonResponse({"error": "No se puede devolver una venta anulada."}, status=400)

    # ── Leer body JSON ──────────────────────────────────────
    try:
        body = json.loads(request.body)
    except json.JSONDecodeError:
        return JsonResponse({"error": "Datos inválidos."}, status=400)

    motivo = body.get("motivo", "").strip()
    items_devolver = body.get("items", [])  # [{detalle_id, cantidad}]

    # ── Validación 2: motivo obligatorio ────────────────────
    if not motivo:
        return JsonResponse({"error": "El motivo de devolución es obligatorio."}, status=400)
    if not _texto_seguro(motivo):
        return JsonResponse({"error": "El motivo solo puede contener letras, números y espacios."}, status=400)

    # ── Validación 3: al menos un ítem ─────────────────────
    if not items_devolver:
        return JsonResponse({"error": "Selecciona al menos un ítem para devolver."}, status=400)

    # ── Validación 4: ítem por ítem ─────────────────────────
    detalles_map = {d.id: d for d in venta.detalles.all()}
    lineas_validadas = []
    total_devolucion = Decimal("0")

    for item in items_devolver:
        detalle_id = item.get("detalle_id")
        try:
            cantidad = int(item.get("cantidad", 0))
        except (TypeError, ValueError):
            return JsonResponse({"error": f"Cantidad inválida en ítem {detalle_id}."}, status=400)

        if cantidad <= 0:
            continue  # el usuario dejó 0, ignorar silenciosamente

        # Existe en esta venta
        if detalle_id not in detalles_map:
            return JsonResponse({"error": f"El ítem {detalle_id} no pertenece a esta venta."}, status=400)

        detalle = detalles_map[detalle_id]
        disponible = detalle.cantidad_disponible_devolver

        # No superar disponible
        if cantidad > disponible:
            nombre = detalle.producto.nombre if detalle.producto else (
                detalle.servicio.nombre if detalle.servicio else f"ítem #{detalle_id}"
            )
            return JsonResponse({
                "error": f"'{nombre}': se intenta devolver {cantidad} pero solo hay {disponible} disponibles para devolución."
            }, status=400)

        subtotal = detalle.precio_unitario * cantidad
        total_devolucion += subtotal
        lineas_validadas.append((detalle, cantidad, subtotal))

    if not lineas_validadas:
        return JsonResponse({"error": "Ingresa al menos una cantidad mayor a 0."}, status=400)

    # ── Crear devolución ────────────────────────────────────
    devolucion = DevolucionVenta.objects.create(
        venta=venta,
        motivo=motivo,
        total_devuelto=total_devolucion,
    )

    for detalle, cantidad, subtotal in lineas_validadas:
        DetalleDevolucion.objects.create(
            devolucion=devolucion,
            detalle_venta=detalle,
            cantidad_devuelta=cantidad,
            subtotal_devuelto=subtotal,
        )
        # Restaurar stock si es producto
        if detalle.producto:
            _ajustar_stock_producto(detalle.producto, cantidad)

    # ── Anulación automática si todos los ítems fueron devueltos ──
    venta_anulada = False
    todos_devueltos = all(
        d.cantidad_disponible_devolver == 0
        for d in venta.detalles.all()
    )
    if todos_devueltos:
        venta.estado = "anulada"
        venta.save(update_fields=["estado"])
        venta_anulada = True

    return JsonResponse({
        "ok": True,
        "codigo_devolucion": devolucion.codigo_devolucion,
        "total_devuelto": float(total_devolucion),
        "venta_anulada": venta_anulada,
        "mensaje": f"Devolución {devolucion.codigo_devolucion} registrada por ${total_devolucion:.2f}.",
    })
