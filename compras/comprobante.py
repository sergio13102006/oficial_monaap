from io import BytesIO

from django.http import HttpResponse
from django.utils.text import slugify
from django.contrib.staticfiles import finders

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as XLImage

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    Image as RLImage,
)

PAGE_WIDTH, PAGE_HEIGHT = A4


def _nombre_archivo_compra(compra, extension):
    proveedor = getattr(compra.proveedor, "nombre_proveedor", "proveedor")
    proveedor_slug = slugify(proveedor) or "proveedor"
    return f"comprobante_compra_{compra.id}_{proveedor_slug}.{extension}"


def _fmt_money(value):
    try:
        return f"${int(value):,}".replace(",", ".")
    except Exception:
        return "$0"


def _find_first_static(paths):
    for path in paths:
        found = finders.find(path)
        if found:
            return found
    return None


def _get_logo_path():
    return _find_first_static([
        "compras/img/logo_monakeratina.png",
        "compras/img/logo_monakeratina.webp",
        "compras/img/logo_monakeratina.jpg",
        "compras/img/logo_monakefratina.png",
        "compras/img/logo_monakefratina.webp",
        "compras/img/logo_monakefratina.jpg",
        "compras/img/logo_monakeratina_watermark.png",
        "compras/img/logo_monakefratina_watermark.png",
        "core/img/logo_monakeratina.png",
        "core/img/logo_monakefratina.png",
        "core/img/logo_monakeratina_watermark.png",
        "core/img/logo_monakefratina_watermark.png",
    ])


def _get_watermark_path():
    return _find_first_static([
        "compras/img/logo_monakeratina_watermark.png",
        "compras/img/logo_monakefratina_watermark.png",
        "compras/img/logo_monakeratina.png",
        "compras/img/logo_monakefratina.png",
        "core/img/logo_monakeratina_watermark.png",
        "core/img/logo_monakefratina_watermark.png",
        "core/img/logo_monakeratina.png",
        "core/img/logo_monakefratina.png",
    ])


def _draw_watermark(canvas, doc):
    watermark_path = _get_watermark_path()
    if not watermark_path:
        return

    canvas.saveState()

    try:
        if hasattr(canvas, "setFillAlpha"):
            canvas.setFillAlpha(0.06)
    except Exception:
        pass

    image_width = 18 * mm
    image_height = 18 * mm
    gap_x = 16 * mm
    gap_y = 18 * mm

    x = 10 * mm
    while x < PAGE_WIDTH:
        y = 12 * mm
        while y < PAGE_HEIGHT:
            canvas.drawImage(
                watermark_path,
                x,
                y,
                width=image_width,
                height=image_height,
                preserveAspectRatio=True,
                mask="auto",
            )
            y += image_height + gap_y
        x += image_width + gap_x

    canvas.restoreState()


def build_comprobante_pdf_response(compra):
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#000000"),
        alignment=1,  # centrado
        spaceAfter=2,
    )

    small_label_style = ParagraphStyle(
        "SmallLabelStyle",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=9,
        textColor=colors.HexColor("#333333"),
    )

    small_value_style = ParagraphStyle(
        "SmallValueStyle",
        parent=styles["BodyText"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=13,
        textColor=colors.HexColor("#111111"),
    )

    text_style = ParagraphStyle(
        "TextStyle",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=9,
        leading=11,
        textColor=colors.HexColor("#222222"),
    )

    footer_style = ParagraphStyle(
        "FooterStyle",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#666666"),
        alignment=1,
    )

    elements = []

    proveedor = getattr(compra.proveedor, "nombre_proveedor", "—")
    usuario = compra.usuario.username if compra.usuario else "—"
    fecha = compra.fecha.strftime("%d/%m/%Y") if compra.fecha else "—"
    estado = "Anulada" if compra.anulada else "Activa"
    total = compra.precio_total or 0

    # Header con logo + título
    logo_path = _get_logo_path()
    logo_flowable = ""
    if logo_path:
        try:
            logo_flowable = RLImage(logo_path, width=34 * mm, height=16 * mm)
        except Exception:
            logo_flowable = ""

    header_table = Table(
        [[logo_flowable, Paragraph("Comprobante", title_style), ""]],
        colWidths=[42 * mm, 100 * mm, 26 * mm],
    )
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, 0), "LEFT"),
        ("ALIGN", (1, 0), (1, 0), "CENTER"),
        ("LINEBELOW", (0, 0), (-1, 0), 1.2, colors.black),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 8))

    # Primera línea tipo factura
    numero_box = Table([
        [Paragraph("NÚMERO", small_label_style)],
        [Paragraph(f"{compra.id:04d}", small_value_style)],
    ], colWidths=[28 * mm])

    de_box = Table([
        [Paragraph("DE", small_label_style)],
        [Paragraph("Monakeratina", small_value_style)],
        [Paragraph("Sistema de compras", text_style)],
        [Paragraph("Documento interno", text_style)],
    ], colWidths=[64 * mm])

    para_box = Table([
        [Paragraph("PARA", small_label_style)],
        [Paragraph(str(proveedor), small_value_style)],
        [Paragraph("Proveedor asociado a la compra", text_style)],
    ], colWidths=[74 * mm])

    top_info = Table(
        [[numero_box, de_box, para_box]],
        colWidths=[30 * mm, 66 * mm, 76 * mm],
    )
    top_info.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEABOVE", (0, 0), (-1, -1), 1.0, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(top_info)
    elements.append(Spacer(1, 10))

    # Segunda línea meta
    fecha_box = Table([
        [Paragraph("FECHA", small_label_style)],
        [Paragraph(fecha, small_value_style)],
    ], colWidths=[28 * mm])

    usuario_box = Table([
        [Paragraph("USUARIO", small_label_style)],
        [Paragraph(usuario, small_value_style)],
    ], colWidths=[28 * mm])

    estado_box = Table([
        [Paragraph("ESTADO", small_label_style)],
        [Paragraph(estado, small_value_style)],
    ], colWidths=[28 * mm])

    meta_table = Table(
        [[fecha_box, usuario_box, estado_box]],
        colWidths=[30 * mm, 30 * mm, 30 * mm],
    )
    meta_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEABOVE", (0, 0), (-1, -1), 1.0, colors.black),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 10))

    # Tabla de productos
    table_data = [["Descripción", "Cantidad", "Precio unidad", "Importe"]]

    detalles = list(compra.detalles.all())
    if detalles:
        for detalle in detalles:
            subtotal = getattr(detalle, "subtotal", 0) or 0
            table_data.append([
                str(detalle.producto.nombre),
                str(detalle.cantidad),
                _fmt_money(detalle.precio_unitario),
                _fmt_money(subtotal),
            ])
    else:
        table_data.append(["No hay productos en esta compra.", "", "", ""])

    detail_table = Table(
        table_data,
        colWidths=[88 * mm, 26 * mm, 32 * mm, 32 * mm],
        repeatRows=1,
    )
    detail_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 1.0, colors.black),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#777777")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("ALIGN", (1, 1), (1, -1), "CENTER"),
        ("ALIGN", (2, 1), (3, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 1), (-1, -1), 0.35, colors.HexColor("#CFCFCF")),
    ]))
    elements.append(detail_table)
    elements.append(Spacer(1, 6))

    # Total abajo a la derecha
    total_table = Table(
        [["Total", _fmt_money(total)]],
        colWidths=[34 * mm, 34 * mm],
        hAlign="RIGHT",
    )
    total_table.setStyle(TableStyle([
        ("LINEABOVE", (0, 0), (-1, 0), 1.0, colors.black),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (0, 0), (0, 0), "LEFT"),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(total_table)

    elements.append(Spacer(1, 10))
    elements.append(Paragraph("Generado por Monakeratina", footer_style))

    doc.build(
        elements,
        onFirstPage=_draw_watermark,
        onLaterPages=_draw_watermark,
    )

    pdf = buffer.getvalue()
    buffer.close()

    filename = _nombre_archivo_compra(compra, "pdf")
    response = HttpResponse(pdf, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response

def build_comprobante_excel_response(compra):
    # =========================
    # Crear libro y hoja
    # =========================
    wb = Workbook()
    ws = wb.active
    ws.title = "Comprobante"

    # =========================
    # Helpers de estilo
    # =========================
    def apply_style(cell, font=None, fill=None, border=None, alignment=None, number_format=None):
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if border:
            cell.border = border
        if alignment:
            cell.alignment = alignment
        if number_format:
            cell.number_format = number_format

    def style_range(start_row, end_row, start_col, end_col, **styles):
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                apply_style(ws.cell(row=row, column=col), **styles)

    # =========================
    # Configuración visual de hoja
    # =========================
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B12"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = False
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4

    # =========================
    # Anchos de columnas y alturas
    # A queda como margen visual
    # =========================
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18

    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 34
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[11].height = 22
    ws.row_dimensions[12].height = 22

    # =========================
    # Paleta visual
    # =========================
    color_dark = "3B261A"
    color_brown = "A67C52"
    color_soft = "F7F0E8"
    color_soft_2 = "FBF7F2"
    color_line = "D8C3B5"
    color_white = "FFFFFF"
    color_text = "2F241D"

    # =========================
    # Fuentes
    # =========================
    title_font = Font(bold=True, size=16, color=color_text)
    section_font = Font(bold=True, size=11, color=color_white)
    label_font = Font(bold=True, size=10, color=color_dark)
    value_font = Font(size=10, color=color_text)
    value_bold_font = Font(bold=True, size=10, color=color_text)
    table_header_font = Font(bold=True, size=10, color=color_white)
    total_font = Font(bold=True, size=11, color=color_dark)
    footer_font = Font(italic=True, size=9, color=color_brown)

    # =========================
    # Rellenos
    # =========================
    dark_fill = PatternFill("solid", fgColor=color_dark)
    brown_fill = PatternFill("solid", fgColor=color_brown)
    soft_fill = PatternFill("solid", fgColor=color_soft)
    soft_fill_2 = PatternFill("solid", fgColor=color_soft_2)

    # =========================
    # Bordes
    # =========================
    thin_side = Side(style="thin", color=color_line)
    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    # =========================
    # Alineaciones
    # =========================
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    right = Alignment(horizontal="right", vertical="center")

    # =========================
    # Datos del comprobante
    # =========================
    proveedor = getattr(compra.proveedor, "nombre_proveedor", "—")
    usuario = compra.usuario.username if compra.usuario else "—"
    fecha = compra.fecha.strftime("%d/%m/%Y") if compra.fecha else "—"
    estado = "Anulada" if compra.anulada else "Activa"
    total = float(compra.precio_total or 0)

    # =========================
    # Logo
    # =========================
    logo_path = _get_logo_path()
    if logo_path:
        try:
            img = XLImage(logo_path)
            img.width = 145
            img.height = 58
            ws.add_image(img, "B1")
        except Exception:
            pass

    # =========================
    # Título principal
    # =========================
    style_range(1, 2, 2, 5, border=thin_border, alignment=center)
    ws.merge_cells("B1:E2")
    ws["B1"] = f"Comprobante de Compra #{compra.id}"
    apply_style(ws["B1"], font=title_font, alignment=center)

    # =========================
    # Sección de información general
    # =========================
    style_range(4, 4, 2, 5, fill=dark_fill, border=thin_border, alignment=left)
    ws.merge_cells("B4:E4")
    ws["B4"] = "DATOS GENERALES"
    apply_style(ws["B4"], font=section_font, alignment=left)

    info_rows = [
        (5, "Proveedor", proveedor),
        (6, "Usuario", usuario),
        (7, "Fecha", fecha),
        (8, "Estado", estado),
        (9, "Total", total),
    ]

    for row_num, label, value in info_rows:
        ws[f"B{row_num}"] = label
        apply_style(ws[f"B{row_num}"], font=label_font, fill=soft_fill, border=thin_border, alignment=left)

        style_range(row_num, row_num, 3, 5, fill=soft_fill_2, border=thin_border, alignment=left)
        ws.merge_cells(start_row=row_num, start_column=3, end_row=row_num, end_column=5)
        ws.cell(row=row_num, column=3, value=value)

        if label == "Total":
            apply_style(
                ws.cell(row=row_num, column=3),
                font=value_bold_font,
                fill=soft_fill_2,
                border=thin_border,
                alignment=right,
                number_format='$#,##0'
            )
        else:
            apply_style(
                ws.cell(row=row_num, column=3),
                font=value_font,
                fill=soft_fill_2,
                border=thin_border,
                alignment=left
            )

    # =========================
    # Sección de detalle
    # =========================
    style_range(11, 11, 2, 5, fill=dark_fill, border=thin_border, alignment=left)
    ws.merge_cells("B11:E11")
    ws["B11"] = "DETALLE DE PRODUCTOS"
    apply_style(ws["B11"], font=section_font, alignment=left)

    headers = ["Producto", "Cantidad", "Precio Unitario", "Subtotal"]
    for col_num, header in enumerate(headers, start=2):
        cell = ws.cell(row=12, column=col_num, value=header)
        apply_style(
            cell,
            font=table_header_font,
            fill=brown_fill,
            border=thin_border,
            alignment=center
        )

    # =========================
    # Filas de productos
    # =========================
    current_row = 13
    detalles = list(compra.detalles.all())

    if detalles:
        for detalle in detalles:
            subtotal = float(getattr(detalle, "subtotal", 0) or 0)

            ws.cell(row=current_row, column=2, value=str(detalle.producto.nombre))
            ws.cell(row=current_row, column=3, value=int(detalle.cantidad or 0))
            ws.cell(row=current_row, column=4, value=float(detalle.precio_unitario or 0))
            ws.cell(row=current_row, column=5, value=subtotal)

            row_fill = soft_fill if current_row % 2 == 1 else soft_fill_2

            for col in range(2, 6):
                apply_style(
                    ws.cell(row=current_row, column=col),
                    fill=row_fill,
                    border=thin_border,
                    alignment=left if col == 2 else center if col == 3 else right
                )

            ws.cell(row=current_row, column=4).number_format = '$#,##0'
            ws.cell(row=current_row, column=5).number_format = '$#,##0'
            current_row += 1
    else:
        ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=5)
        ws.cell(row=current_row, column=2, value="No hay productos en esta compra.")
        style_range(current_row, current_row, 2, 5, fill=soft_fill_2, border=thin_border, alignment=center)
        apply_style(ws.cell(row=current_row, column=2), font=value_font, alignment=center)
        current_row += 1

    # =========================
    # Total final
    # =========================
    total_row = current_row + 1

    apply_style(
        ws.cell(row=total_row, column=4, value="TOTAL"),
        font=total_font,
        fill=soft_fill,
        border=thin_border,
        alignment=right
    )
    apply_style(
        ws.cell(row=total_row, column=5, value=total),
        font=total_font,
        fill=soft_fill,
        border=thin_border,
        alignment=right,
        number_format='$#,##0'
    )

    # =========================
    # Footer
    # =========================
    footer_row = total_row + 2
    style_range(footer_row, footer_row, 2, 5, alignment=center)
    ws.merge_cells(start_row=footer_row, start_column=2, end_row=footer_row, end_column=5)
    ws.cell(row=footer_row, column=2, value="Generado por MonaApp / Monakeratina")
    apply_style(ws.cell(row=footer_row, column=2), font=footer_font, alignment=center)

    # =========================
    # Área de impresión
    # =========================
    ws.print_area = f"B1:E{footer_row}"

    # =========================
    # Respuesta HTTP
    # =========================
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = _nombre_archivo_compra(compra, "xlsx")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response